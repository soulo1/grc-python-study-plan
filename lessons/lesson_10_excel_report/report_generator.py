"""
Lesson 10 - Compliance Report Generator (Excel)
==============================================

GRC problem: Leadership and auditors want a clean control-status report, not a
raw CSV. Building it by hand every month is tedious.

This tool reads control status data and produces a polished .xlsx workbook with
a color-coded summary sheet and a detailed per-control sheet.

This is your first use of a THIRD-PARTY library, openpyxl. Install it with:
    pip install openpyxl
(or `pip install -r ../../requirements.txt` from the project).

Run it:
    python report_generator.py            # reads controls_status.csv
    -> writes compliance_report.xlsx

Concepts introduced:
- installing and importing a third-party package (openpyxl)
- creating workbooks/worksheets, writing cells
- styling: fonts, fills, column widths
- conditional (red/amber/green) formatting driven by data
"""

import csv
import sys
from collections import Counter
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("This lesson needs openpyxl. Install it with:\n    pip install openpyxl")
    sys.exit(1)

# Map a status to a fill color (RAG: red/amber/green).
STATUS_FILL = {
    "Effective": "C6EFCE",          # green
    "Needs Improvement": "FFEB9C",  # amber
    "Not Tested": "FFC7CE",         # red
    "Ineffective": "FFC7CE",        # red
}
STATUS_FONT = {
    "Effective": "006100",
    "Needs Improvement": "9C6500",
    "Not Tested": "9C0006",
    "Ineffective": "9C0006",
}

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def load_controls(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def build_summary_sheet(ws, controls):
    ws.title = "Summary"
    counts = Counter(c["status"] for c in controls)
    total = len(controls)

    ws["A1"] = "Compliance Control Status Summary"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = f"Total controls: {total}"

    ws.append([])  # blank row
    ws.append(["Status", "Count", "Percent"])
    header_row = ws.max_row
    for col in range(1, 4):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for status, count in counts.most_common():
        pct = f"{count / total * 100:.0f}%"
        ws.append([status, count, pct])
        row = ws.max_row
        fill = STATUS_FILL.get(status)
        if fill:
            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10


def build_detail_sheet(wb, controls):
    ws = wb.create_sheet("Controls")
    headers = ["Control ID", "Name", "Framework", "Owner", "Status", "Last Tested"]
    ws.append(headers)
    style_header(ws, len(headers))

    for c in controls:
        ws.append([
            c["control_id"], c["name"], c["framework"],
            c["owner"], c["status"], c.get("last_tested", "") or "-",
        ])
        row = ws.max_row
        status = c["status"]
        fill = STATUS_FILL.get(status)
        if fill:
            cell = ws.cell(row=row, column=5)  # the Status column
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(color=STATUS_FONT.get(status, "000000"), bold=True)

    widths = [12, 34, 12, 16, 18, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"  # keep header visible when scrolling


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "controls_status.csv"
    out = "compliance_report.xlsx"

    controls = load_controls(src)
    wb = Workbook()
    build_summary_sheet(wb.active, controls)
    build_detail_sheet(wb, controls)
    wb.save(out)

    print("Compliance Report Generator")
    print("=" * 50)
    print(f"Read {len(controls)} controls from {src}")
    print(f"Wrote {Path(out).resolve()}")
    print("Open it in Excel/LibreOffice/Google Sheets - the Status column is")
    print("color-coded red/amber/green for instant readability.")


if __name__ == "__main__":
    main()
